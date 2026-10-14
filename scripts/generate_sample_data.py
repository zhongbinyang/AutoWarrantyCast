"""Generate realistic synthetic sample data for the warranty cost system.

Produces ``sample_data/保修预测导入示例.xlsx`` with instructions, monthly
sales, mileage, two systems' historical triangles and held-out actuals.
Older separate sample files are retained as regression fixtures.

Each system uses a distinct true Weibull hazard so the two triangles look
meaningfully different, with Poisson claim-count noise and lognormal cost
noise layered on top. Cells beyond each batch's observed development window
(governed by a single fixed "data evaluation month") are left blank, exactly
matching the ragged-triangle shape of real warranty data.

Run with: ``python scripts/generate_sample_data.py``
"""

from __future__ import annotations

import os
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

RNG = np.random.default_rng(42)
OUT_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "sample_data")
DEMO_NAME = "保修预测导入示例.xlsx"

N_BATCHES = 30
START_MONTH = "2022-01"
EVAL_MONTH = "2024-07"  # data snapshot date

SYSTEMS = {
    "engine": {
        "sheet_label": "发动机系统",
        "beta_t": 2.3,
        "eta_t": 55.0,
        "base_unit_cost": 3400.0,
        "cost_growth": 0.01,
        # Engine wear is genuinely mileage-driven in this sample, so the
        # bivariate time+mileage model has real signal to recover.
        "beta_m": 1.8,
        "eta_m": 45000.0,
    },
    "chassis": {
        "sheet_label": "底盘系统",
        "beta_t": 1.6,
        "eta_t": 58.0,
        "base_unit_cost": 1800.0,
        "cost_growth": 0.006,
        # Chassis failures are treated as calendar-time driven only.
        "beta_m": None,
        "eta_m": None,
    },
}

MAX_AGE = 36


def month_range(start: str, n: int):
    p = pd.Period(start, freq="M")
    return [str(p + i) for i in range(n)]


def build_sales(rng=RNG) -> pd.DataFrame:
    months = month_range(START_MONTH, N_BATCHES)
    base = 4000
    trend = np.linspace(0, 2500, N_BATCHES)
    seasonal = 800 * np.sin(np.linspace(0, 6 * np.pi, N_BATCHES))
    noise = rng.normal(0, 250, N_BATCHES)
    sales = np.clip(base + trend + seasonal + noise, 500, None).round().astype(int)
    return pd.DataFrame({"销售月份": months, "销量": sales})


def observed_age_limit(sale_month: str, evaluation_month: str = EVAL_MONTH) -> int:
    age = (pd.Period(evaluation_month, freq="M") - pd.Period(sale_month, freq="M")).n
    return int(np.clip(age, 0, MAX_AGE))


def build_system_triangles(
    sales_df: pd.DataFrame, params: dict, mileage_df: pd.DataFrame = None,
    evaluation_month: str = EVAL_MONTH, rng=RNG,
):
    beta_t = params["beta_t"]
    eta_t = params["eta_t"]
    base_cost = params["base_unit_cost"]
    cost_growth = params["cost_growth"]
    beta_m = params.get("beta_m")
    eta_m = params.get("eta_m")

    mileage_map = {}
    if mileage_df is not None:
        mileage_map = dict(zip(mileage_df["销售月份"], mileage_df["月均行驶里程"]))

    ages = list(range(0, MAX_AGE + 1))
    count_rows = []
    amount_rows = []

    for _, row in sales_df.iterrows():
        sm = row["销售月份"]
        sales = int(row["销量"])
        max_obs_age = observed_age_limit(sm, evaluation_month)
        batch_factor = float(np.exp(rng.normal(0, 0.08)))
        monthly_mileage = float(mileage_map.get(sm, 0.0))

        cum_prev = 0.0
        count_row = {"销售月份": sm}
        amount_row = {"销售月份": sm}
        for age in ages:
            col = f"MIS{age}"
            if age > max_obs_age:
                count_row[col] = np.nan
                amount_row[col] = np.nan
                continue
            hazard = (age / eta_t) ** beta_t
            if beta_m is not None and eta_m is not None and monthly_mileage > 0:
                cumulative_mileage = age * monthly_mileage
                hazard += (cumulative_mileage / eta_m) ** beta_m
            true_cum_rate = 1.0 - np.exp(-hazard)
            true_cum_rate *= batch_factor
            true_cum_rate = min(true_cum_rate, 0.98)
            expected_cum_count = true_cum_rate * sales
            expected_inc = max(expected_cum_count - cum_prev, 0.0)
            inc_count = rng.poisson(lam=expected_inc) if expected_inc > 0 else 0
            cum_prev += inc_count

            unit_cost = base_cost * (1 + cost_growth * age) * float(np.exp(rng.normal(0, 0.07)))
            inc_amount = inc_count * unit_cost * float(np.exp(rng.normal(0, 0.05)))

            count_row[col] = int(inc_count)
            amount_row[col] = round(float(inc_amount), 2)

        count_rows.append(count_row)
        amount_rows.append(amount_row)

    return pd.DataFrame(count_rows), pd.DataFrame(amount_rows)


def build_mileage(sales_df: pd.DataFrame, rng=RNG) -> pd.DataFrame:
    n = len(sales_df)
    base = 1500
    drift = np.linspace(-150, 150, n)
    noise = rng.normal(0, 120, n)
    mileage = np.clip(base + drift + noise, 800, None).round().astype(int)
    return pd.DataFrame({"销售月份": sales_df["销售月份"], "月均行驶里程": mileage})


def build_demo_tables() -> dict[str, pd.DataFrame]:
    """Generate one coherent snapshot plus three held-out validation months."""
    rng = np.random.default_rng(42)
    sales = build_sales(rng)
    mileage = build_mileage(sales, rng)
    tables = {"销量": sales, "月均里程": mileage}
    for params in SYSTEMS.values():
        count, amount = build_system_triangles(sales, params, mileage, "2024-10", rng)
        for metric, full in [("台次", count), ("金额", amount)]:
            historical, actual = full.copy(), full.copy()
            for row, month in enumerate(sales["销售月份"]):
                last_observed = observed_age_limit(month)
                historical.loc[row, [f"MIS{age}" for age in range(last_observed + 1, MAX_AGE + 1)]] = np.nan
                actual.loc[row, [f"MIS{age}" for age in range(last_observed + 1)]] = np.nan
            system = params["sheet_label"]
            tables[f"{system}_{metric}"] = historical
            tables[f"{system}_验证{metric}"] = actual
    return tables


def _instructions_sheet(workbook: Workbook) -> None:
    sheet = workbook.active
    sheet.title = "使用说明"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.tabColor = "17324D"
    for col in "ABCDEF":
        sheet.column_dimensions[col].width = 20

    def line(row, text, *, heading=False, title=False):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        cell = sheet.cell(row, 1, text)
        cell.font = Font(name="Microsoft YaHei", size=20 if title else 12 if heading else 11,
                         bold=title or heading, color="FFFFFF" if title or heading else "243B53")
        cell.fill = PatternFill("solid", fgColor="17324D" if title else "237E85" if heading else "F3F7FA")
        cell.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
        sheet.row_dimensions[row].height = 44 if title else 28 if heading else 34

    line(1, "车辆保修预测 · Excel 导入示例", title=True)
    line(2, "模拟数据 / 固定种子 42 / 一次上传全部数据 / 金额单位：元")
    line(4, "01   三步开始", heading=True)
    line(5, "① 上传本文件 → ② 检查各 Sheet 的自动识别结果，点击“导入数据” → ③ 点击“运行预测模型”。")
    line(6, "本文件含 30 个销售批次、发动机和底盘两个系统；可直接导入，无需拆分文件。")
    line(8, "02   Sheet 与字段", heading=True)
    line(9, "销量（必填）：销售月份、销量。月份采用 YYYY-MM，销量为非负整数，总销量必须大于 0。")
    line(10, "系统名_台次 / 系统名_金额（必填）：销售月份、MIS0、MIS1……；每个系统必须成对提供。")
    line(11, "月均里程（可选）：销售月份、月均行驶里程（公里/月）。每个月份一行；导入后默认使用。")
    line(12, "系统名_验证台次 / 系统名_验证金额（可选）：实际增量值，可只提供其中一项；仅用于 MAE / MAPE。")
    line(14, "03   填写与识别规则", heading=True)
    line(15, "数据 Sheet 第一行必须为表头；保留销售月份和 MIS 列名，不要在数据区添加合计行或说明。")
    line(16, "已知没有索赔填 0；未知或未发生值留空。台次为非负整数，金额为非负数值，单位统一为元。")
    line(17, "历史数据默认采用增量；系统保留累积口径自动识别和转换。验证数据必须为增量，不会进入训练。")
    line(18, "Sheet 名称不规范时，可在页面修改用途、系统和指标；无需导入的 Sheet 设为“忽略”。")
    line(19, "导入失败会标明 Sheet 和问题；修改文件或映射后须重新导入，原预测和下载结果将失效。")
    line(21, "04   示例时间范围与颜色", heading=True)
    line(22, "销售月份：2022-01 至 2024-06；历史截止：2024-07；验证区间：2024-08 至 2024-10。")
    line(23, "蓝绿色表头：销量、里程和历史数据；紫色表头：验证数据；浅灰空白：未知或不参与该用途的区域。")
    line(24, "来源：scripts/generate_sample_data.py 的可复现模拟数据；不代表真实车辆、真实索赔或精算建议。")
    line(26, "05   支持的系统", heading=True)
    line(27, "发动机系统、底盘系统、传动系统、转向系统、车身系统、车身内饰、车门天窗、电气系统、三电系统。")
    sheet.freeze_panes = "A4"
    sheet.print_area = "A1:F27"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1


def write_demo_workbook(path: Path | str) -> Path:
    """Write the import-ready, formatted demo using the project's openpyxl dependency."""
    tables = build_demo_tables()
    workbook = Workbook()
    workbook.properties.title = "车辆保修预测统一 Excel 导入示例"
    workbook.properties.subject = "销量、历史索赔、里程与验证数据的一次性导入"
    workbook.properties.creator = "AutoWarrantyCast"
    _instructions_sheet(workbook)
    # Keep historical pairs together and place held-out data at the end.
    ordered = [name for name in tables if "验证" not in name] + [name for name in tables if "验证" in name]
    for name in ordered:
        frame = tables[name]
        sheet = workbook.create_sheet(name)
        sheet.sheet_view.showGridLines = False
        sheet.sheet_view.zoomScale = 90
        validation = "验证" in name
        color = "73509B" if validation else "237E85"
        sheet.sheet_properties.tabColor = color
        sheet.append(list(frame.columns))
        for row in frame.itertuples(index=False, name=None):
            sheet.append([None if pd.isna(value) else value for value in row])
        sheet.freeze_panes = "B2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 30
        sheet.column_dimensions["A"].width = 16
        for index in range(2, len(frame.columns) + 1):
            sheet.column_dimensions[get_column_letter(index)].width = 19 if "金额" in name or len(frame.columns) == 2 else 13
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor=color)
            cell.font = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="left" if cell.column == 1 else "right", vertical="center")
            cell.border = Border(bottom=Side(style="thin", color="D5E2EA"))
        for cells in sheet.iter_rows(min_row=2):
            sheet.row_dimensions[cells[0].row].height = 23
            for cell in cells:
                cell.font = Font(name="Microsoft YaHei", size=11, color="243B53")
                cell.alignment = Alignment(horizontal="left" if cell.column == 1 else "right", vertical="center")
                cell.fill = PatternFill("solid", fgColor="EEF1F4" if cell.value is None else "F3F8FA" if cell.row % 2 == 0 else "FFFFFF")
                cell.number_format = "@" if cell.column == 1 else "#,##0.00" if "金额" in name else "#,##0"
        note = "增量实际数据；仅包含 2024-08 至 2024-10，供预测验证，不参与训练。" if validation else "历史增量数据；截至 2024-07，未来未知单元格留空。"
        if name == "销量":
            note = "月度销量，非负整数；固定种子 42 生成的模拟数据。"
        elif name == "月均里程":
            note = "单位：公里/月；每个销售月份一条记录。"
        sheet["A1"].comment = Comment(note, "AutoWarrantyCast")
        sheet.print_title_rows = "1:1"
        sheet.print_title_cols = "A:A"
        sheet.print_area = sheet.dimensions
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        # Four or five horizontal pages for wide triangles keeps MIS columns
        # readable on paper without changing the machine-readable data grid.
        sheet.page_setup.fitToWidth = 0 if len(frame.columns) > 2 else 1
        sheet.page_setup.fitToHeight = 1
        sheet.oddFooter.center.text = f"{name} | 模拟数据 | &P / &N"
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()
    return path


def main() -> None:
    path = write_demo_workbook(Path(OUT_DIR) / DEMO_NAME)
    print(f"统一 Excel Demo 已生成：{path}")


if __name__ == "__main__":
    main()
