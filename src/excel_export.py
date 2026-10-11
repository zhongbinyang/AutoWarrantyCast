"""Excel/CSV export helpers, including light-yellow highlighting of forecast cells."""

from __future__ import annotations

from io import BytesIO
from typing import Dict, List

import pandas as pd
from openpyxl.styles import PatternFill

from .utils import FORECAST_FILL_HEX, SALE_MONTH_COL, SALES_COL, mis_col


def build_derived_triangles(
    count_filled: pd.DataFrame,
    amount_filled: pd.DataFrame,
    sales_df: pd.DataFrame,
    ages: List[int],
) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Derive cumulative-to-date failure-rate and unit-cost-per-vehicle triangles."""
    cols = [mis_col(a) for a in ages]
    sales_map = dict(zip(sales_df[SALE_MONTH_COL], sales_df[SALES_COL]))
    sales_series = pd.Series(sales_map).reindex(count_filled.index).replace(0, pd.NA)

    cum_count = count_filled[cols].fillna(0).cumsum(axis=1)
    cum_count = cum_count.where(count_filled[cols].notna())
    cum_amount = amount_filled[cols].fillna(0).cumsum(axis=1)
    cum_amount = cum_amount.where(amount_filled[cols].notna())

    fail_rate = cum_count.div(sales_series, axis=0)
    unit_cost_per_vehicle = cum_amount.div(sales_series, axis=0)
    return fail_rate, unit_cost_per_vehicle


def export_triangles_workbook(
    sheets: Dict[str, pd.DataFrame], forecast_masks: Dict[str, pd.DataFrame]
) -> bytes:
    """Write an in-memory .xlsx workbook, one sheet per (name -> DataFrame) pair.

    Cells flagged True in the matching ``forecast_masks[name]`` frame are
    highlighted with a light-yellow fill so users can immediately see which
    values are forecast versus actual.
    """
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for name, df in sheets.items():
            out = df.reset_index().rename(columns={"index": SALE_MONTH_COL})
            # Excel sheet names are capped at 31 characters.
            safe_name = name[:31]
            out.to_excel(writer, sheet_name=safe_name, index=False)

    buffer.seek(0)
    import openpyxl

    wb = openpyxl.load_workbook(buffer)
    fill = PatternFill(start_color=FORECAST_FILL_HEX, end_color=FORECAST_FILL_HEX, fill_type="solid")

    for name, df in sheets.items():
        safe_name = name[:31]
        ws = wb[safe_name]
        mask = forecast_masks.get(name)
        if mask is None:
            continue
        header = [cell.value for cell in ws[1]]
        col_idx = {col_name: idx + 1 for idx, col_name in enumerate(header)}
        for row_offset, sale_month in enumerate(df.index, start=2):
            if sale_month not in mask.index:
                continue
            for col_name in df.columns:
                if col_name not in col_idx or col_name not in mask.columns:
                    continue
                if bool(mask.loc[sale_month, col_name]):
                    ws.cell(row=row_offset, column=col_idx[col_name]).fill = fill

    out_buffer = BytesIO()
    wb.save(out_buffer)
    return out_buffer.getvalue()


def export_summary_csv(summary_df: pd.DataFrame) -> bytes:
    """Serialize a summary DataFrame to CSV bytes (UTF-8 with BOM for Excel)."""
    return summary_df.to_csv(index=False).encode("utf-8-sig")
