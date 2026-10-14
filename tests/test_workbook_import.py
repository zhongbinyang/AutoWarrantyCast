"""Unified import regression tests; run with unittest discovery."""

from dataclasses import replace
from pathlib import Path
import unittest

import numpy as np
import pandas as pd
from pandas.testing import assert_frame_equal
from openpyxl import load_workbook

from src import data_io
from scripts.generate_sample_data import build_demo_tables

ROOT = Path(__file__).resolve().parents[1]


def input_sheets():
    months = ["2024-01", "2024-02"]
    return {
        "销量": pd.DataFrame({"销售月份": months, "销量": [100, 200]}),
        "发动机系统_台次": pd.DataFrame({"销售月份": months, "MIS0": [3, 2], "MIS1": [1, np.nan]}),
        "发动机系统_金额": pd.DataFrame({"销售月份": months, "MIS0": [300, 200], "MIS1": [100, np.nan]}),
    }


def load(sheets):
    return data_io.load_workbook_data(sheets, data_io.suggest_sheet_mappings(sheets))


class WorkbookImportTests(unittest.TestCase):
    def test_demo_file_matches_generator_and_preserves_numeric_cells(self):
        path = ROOT / "sample_data" / "保修预测导入示例.xlsx"
        sheets = data_io.read_workbook(path)
        expected = build_demo_tables()
        self.assertEqual(set(sheets), {"使用说明", *expected})
        for name, frame in expected.items():
            assert_frame_equal(sheets[name], frame, check_dtype=False)
        imported = load(sheets)
        self.assertEqual(len(imported.bundles), 2)
        self.assertEqual(len(imported.sales), 30)
        with path.open("rb") as source:
            workbook = load_workbook(source)
            for name in expected:
                sheet = workbook[name]
                self.assertEqual(sheet.freeze_panes, "B2")
                self.assertEqual(sheet["A1"].value, "销售月份")
                for row in sheet.iter_rows(min_row=2, min_col=2):
                    for cell in row:
                        if cell.value is not None:
                            self.assertEqual(cell.data_type, "n")
            self.assertIsNone(workbook["发动机系统_台次"]["AL31"].value)
            workbook.close()

    def test_demo_data_is_reproducible_and_validation_is_held_out(self):
        first, second = build_demo_tables(), build_demo_tables()
        for sheet in first:
            assert_frame_equal(first[sheet], second[sheet])
        result = load(first)
        self.assertEqual(len(result.sales), 30)
        self.assertEqual(len(result.bundles), 2)
        for system, bundle in result.bundles.items():
            self.assertEqual(bundle.evaluation_month, "2024-07")
            for metric in data_io.CLAIM_METRICS:
                actual = result.actuals[system][metric]
                self.assertEqual(int(actual.notna().to_numpy().sum()), 90)
                self.assertFalse((actual.notna() & bundle.observed_mask).to_numpy().any())

    def test_minimal_workbook_and_old_bundle_equivalence(self):
        sheets = input_sheets()
        result = load(sheets)
        old = data_io.prepare_triangle_bundle(
            "发动机系统", sheets["发动机系统_台次"].set_index("销售月份"),
            sheets["发动机系统_金额"].set_index("销售月份"), sheets["销量"],
        )
        new = result.bundles["发动机系统"]
        assert_frame_equal(new.count, old.count)
        assert_frame_equal(new.amount, old.amount)
        assert_frame_equal(new.observed_mask, old.observed_mask)
        self.assertEqual(new.evaluation_month, "2024-02")
        self.assertIsNone(result.mileage)
        self.assertEqual(result.actuals, {})

    def test_all_systems_unordered_and_aliases(self):
        sheets = input_sheets()
        count = sheets.pop("发动机系统_台次")
        amount = sheets.pop("发动机系统_金额")
        sheets["说明"] = pd.DataFrame({"内容": ["demo"]})
        sheets["备注"] = pd.DataFrame({"内容": ["ignore"]})
        for system in data_io.SYSTEM_LIST:
            sheets[f"{system} - 索赔金额"] = amount
            sheets[f"{system}（索赔台次）"] = count
        result = load(dict(reversed(list(sheets.items()))))
        self.assertEqual(set(result.bundles), set(data_io.SYSTEM_LIST))

    def test_manual_mapping_for_unknown_names(self):
        sheets = dict(zip(["data1", "data2", "data3"], input_sheets().values()))
        suggestions = data_io.suggest_sheet_mappings(sheets)
        self.assertEqual([m.purpose for m in suggestions], ["销量", "忽略", "忽略"])
        mappings = [suggestions[0], data_io.SheetMapping("data2", "历史索赔", "发动机系统", "台次"),
                    data_io.SheetMapping("data3", "历史索赔", "发动机系统", "金额")]
        self.assertEqual(len(data_io.load_workbook_data(sheets, mappings).bundles), 1)

    def test_validation_is_incremental_and_never_extends_training(self):
        sheets = input_sheets()
        baseline = load(sheets).bundles["发动机系统"]
        actual = pd.DataFrame({"销售月份": ["2024-01", "2024-02"], "MIS0": [3, 2], "MIS1": [6, 4], "MIS2": [9, 8]})
        sheets["发动机系统_验证台次"] = actual
        result = load(sheets)
        self.assertEqual(result.bundles["发动机系统"].evaluation_month, baseline.evaluation_month)
        assert_frame_equal(result.bundles["发动机系统"].count, baseline.count)
        self.assertEqual(result.actuals["发动机系统"]["台次"].loc["2024-01", "MIS2"], 9)
        self.assertEqual(list(result.actuals["发动机系统"]), ["台次"])

    def test_cumulative_history_keeps_existing_conversion(self):
        sheets = input_sheets()
        sheets["发动机系统_台次"]["MIS1"] = [5, np.nan]
        sheets["发动机系统_金额"]["MIS1"] = [500, np.nan]
        result = load(sheets).bundles["发动机系统"]
        self.assertEqual(result.count.loc["2024-01", "MIS1"], 2)
        self.assertEqual(result.amount.loc["2024-01", "MIS1"], 200)
        self.assertEqual(len(result.issues), 2)

    def test_missing_and_duplicate_mappings(self):
        for missing in ["销量", "发动机系统_台次", "发动机系统_金额"]:
            with self.subTest(missing=missing):
                sheets = input_sheets()
                del sheets[missing]
                with self.assertRaises(ValueError):
                    load(sheets)
        sheets = input_sheets()
        for target in ["销量", "发动机系统_台次"]:
            with self.subTest(duplicate=target):
                sheets2 = {**sheets, target + "_副本": sheets[target]}
                with self.assertRaisesRegex(ValueError, "映射重复"):
                    load(sheets2)
        mappings = data_io.suggest_sheet_mappings(sheets)
        with self.assertRaisesRegex(ValueError, "重复使用"):
            data_io.load_workbook_data(sheets, mappings + [mappings[0]])
        with self.assertRaisesRegex(ValueError, "不存在"):
            data_io.load_workbook_data(sheets, [replace(mappings[0], sheet_name="missing")] + mappings[1:])

    def test_numeric_and_schema_errors_identify_sheet(self):
        for field, value in [("MIS0", -1), ("MIS0", 1.5), ("MIS0", np.inf), ("MIS0", "abc"), ("销售月份", "invalid")]:
            with self.subTest(field=field, value=value):
                sheets = input_sheets()
                sheets["发动机系统_台次"] = sheets["发动机系统_台次"].astype(object)
                sheets["发动机系统_台次"].loc[0, field] = value
                with self.assertRaisesRegex(ValueError, "Sheet「发动机系统_台次」.*第 2 行"):
                    load(sheets)
        for bad in [pd.DataFrame(), pd.DataFrame({"销售月份": ["2024-01"], "其他": [0]}),
                    pd.DataFrame({"销售月份": ["2024-01"], "MIS0": [None]})]:
            sheets = input_sheets()
            sheets["发动机系统_台次"] = bad
            with self.assertRaisesRegex(ValueError, "发动机系统_台次"):
                load(sheets)

    def test_sales_empty_invalid_and_no_claim_system(self):
        for value in [None, "bad", -1, 0.5, float("inf")]:
            with self.subTest(value=value):
                sheets = input_sheets()
                sheets["销量"]["销量"] = [value, 200]
                with self.assertRaisesRegex(ValueError, "销量"):
                    load(sheets)
        sheets = input_sheets()
        sheets["销量"]["销量"] = 0
        with self.assertRaisesRegex(ValueError, "总销量"):
            load(sheets)
        with self.assertRaisesRegex(ValueError, "至少选择一个系统"):
            load({"销量": input_sheets()["销量"]})

    def test_mileage_optional_duplicate_and_missing_month_warning(self):
        sheets = input_sheets()
        sheets["月均里程"] = pd.DataFrame({"销售月份": ["2024-01"], "月均行驶里程": [1500]})
        result = load(sheets)
        self.assertEqual(result.mileage.iloc[0]["月均行驶里程"], 1500)
        self.assertIn("2024-02", result.issues[0].message)
        sheets["月均里程"] = pd.concat([sheets["月均里程"]] * 2, ignore_index=True)
        with self.assertRaisesRegex(ValueError, "多条里程"):
            load(sheets)

    def test_dates_headers_duplicate_rows_and_future_blanks(self):
        sheets = input_sheets()
        for name, df in sheets.items():
            df["销售月份"] = pd.to_datetime(df["销售月份"])
            df.rename(columns={"销售月份": " 销售月份 ", "MIS0": " mis0 "}, inplace=True)
        result = load(sheets)
        self.assertEqual(result.sales.iloc[0]["销售月份"], "2024-01")
        self.assertTrue(pd.isna(result.bundles["发动机系统"].count.loc["2024-02", "MIS1"]))
        sheets = input_sheets()
        sheets["发动机系统_台次"] = pd.concat([sheets["发动机系统_台次"]] * 2, ignore_index=True)
        self.assertEqual(load(sheets).bundles["发动机系统"].count.loc["2024-01", "MIS0"], 6)
        sheets["发动机系统_台次"]["mis0"] = 0
        with self.assertRaisesRegex(ValueError, "重复的 MIS"):
            load(sheets)

    def test_validation_requires_imported_system(self):
        sheets = input_sheets()
        sheets["底盘系统_验证金额"] = sheets["发动机系统_金额"]
        with self.assertRaisesRegex(ValueError, "未导入历史索赔"):
            load(sheets)

    def test_existing_excel_reader_preserves_sheet(self):
        path = ROOT / "sample_data" / "保修预测导入示例.xlsx"
        sheet_name = "发动机系统_台次"
        workbook = data_io.read_workbook(path.read_bytes())
        normalized = data_io._normalize_triangle(workbook[sheet_name], "台次")
        assert_frame_equal(normalized, data_io.load_triangle_sheet(path, sheet_name=sheet_name))
        with self.assertRaises(Exception):
            data_io.read_workbook(b"not an xlsx")


if __name__ == "__main__":
    unittest.main()
